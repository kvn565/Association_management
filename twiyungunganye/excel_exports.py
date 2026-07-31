from django.http import HttpResponse
from django.contrib.auth.decorators import login_required, user_passes_test
from django.db.models import Sum, Q
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter
from decimal import Decimal
from datetime import datetime
from io import BytesIO

from .models import Membre, Cotisation, CotisationImprevu, Credit, Interet, Remboursement, Association, PenaliteCotisation, Suggestion
from cas_assistance.models import CasAssistance


def is_admin_or_gestionnaire(user):
    return user.is_authenticated and (user.is_admin or user.is_admin)


def get_tenant(request):
    return getattr(request, 'tenant', None)


def qs_for_tenant(request, qs, field='tenant'):
    tenant = get_tenant(request)
    if tenant and not getattr(request, 'is_superadmin', False):
        return qs.filter(**{field: tenant})
    return qs


def get_devise(request):
    tenant = get_tenant(request)
    assoc = Association.get_association(tenant)
    return assoc.devise if assoc and assoc.devise else 'FBU'


def _style_header(ws, num_cols):
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4a90e2', end_color='4a90e2', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border


def _style_data(ws, num_rows, num_cols, devise='FBU'):
    data_font = Font(name='Calibri', size=10)
    data_alignment = Alignment(horizontal='center', vertical='center')
    money_alignment = Alignment(horizontal='right', vertical='center')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    for row in range(2, num_rows + 1):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border


def _auto_width(ws, num_cols, max_width=40):
    for col in range(1, num_cols + 1):
        max_len = 0
        letter = get_column_letter(col)
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=False):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 3, max_width)


def _format_money(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    cell.value = float(value)
    cell.number_format = '#,##0.000'


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_membres_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Membres"
    devise = get_devise(request)

    headers = ['#', 'Nom', 'Prénom', 'Carte ID', 'Email', 'Téléphone', 'Adresse', 'Date inscription', 'Actif',
               f'Total cotisations ({devise})', f'Total imprévus ({devise})']
    ws.append(headers)
    _style_header(ws, len(headers))

    membres = qs_for_tenant(request, Membre.objects.all())
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')
    if date_debut:
        membres = membres.filter(date_inscription__gte=date_debut)
    if date_fin:
        membres = membres.filter(date_inscription__lte=date_fin)
    
    membres = membres.order_by('nom', 'prenom')
    for i, m in enumerate(membres, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=m.nom)
        ws.cell(row=row, column=3, value=m.prenom)
        ws.cell(row=row, column=4, value=m.carte_identite)
        ws.cell(row=row, column=5, value=m.email or '')
        ws.cell(row=row, column=6, value=m.telephone or '')
        ws.cell(row=row, column=7, value=m.adresse or '')
        ws.cell(row=row, column=8, value=m.date_inscription.strftime('%d/%m/%Y') if m.date_inscription else '')
        ws.cell(row=row, column=9, value='Oui' if m.est_actif else 'Non')
        _format_money(ws, row, 10, m.get_total_cotisations())
        _format_money(ws, row, 11, m.get_total_cotisations_imprevus())

    _style_data(ws, len(membres) + 1, len(headers), devise)
    _auto_width(ws, len(headers))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Membres.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_cotisations_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cotisations"
    devise = get_devise(request)

    headers = ['#', 'Membre', 'Carte ID', 'Montant', 'Mois', 'Année', 'Statut', 'Date paiement']
    ws.append(headers)
    _style_header(ws, len(headers))

    cotisations = qs_for_tenant(request, Cotisation.objects.all()).select_related('membre').order_by('-annee', '-mois')
    mois_noms = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    total = Decimal('0')
    for i, c in enumerate(cotisations, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=c.membre.get_full_name())
        ws.cell(row=row, column=3, value=c.membre.carte_identite)
        _format_money(ws, row, 4, c.montant)
        ws.cell(row=row, column=5, value=mois_noms[c.mois])
        ws.cell(row=row, column=6, value=c.annee)
        ws.cell(row=row, column=7, value=c.get_statut_display())
        ws.cell(row=row, column=8, value=c.date_paiement.strftime('%d/%m/%Y') if c.date_paiement else '')
        if c.statut == 'PAYE':
            total += c.montant

    total_row = len(cotisations) + 2
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    _format_money(ws, total_row, 4, total)
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    _style_data(ws, total_row, len(headers), devise)
    _auto_width(ws, len(headers))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Cotisations.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_cotisations_imprevus_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cotisations Imprévus"
    devise = get_devise(request)

    headers = ['#', 'Membre', 'Carte ID', 'Montant', 'Mois', 'Année', 'Statut', 'Date paiement']
    ws.append(headers)
    _style_header(ws, len(headers))

    cotisations = qs_for_tenant(request, CotisationImprevu.objects.all()).select_related('membre').order_by('-annee', '-mois')
    mois_noms = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    total = Decimal('0')
    for i, c in enumerate(cotisations, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=c.membre.get_full_name())
        ws.cell(row=row, column=3, value=c.membre.carte_identite)
        _format_money(ws, row, 4, c.montant)
        ws.cell(row=row, column=5, value=mois_noms[c.mois])
        ws.cell(row=row, column=6, value=c.annee)
        ws.cell(row=row, column=7, value=c.get_statut_display())
        ws.cell(row=row, column=8, value=c.date_paiement.strftime('%d/%m/%Y') if c.date_paiement else '')
        if c.statut == 'PAYE':
            total += c.montant

    total_row = len(cotisations) + 2
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    _format_money(ws, total_row, 4, total)
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    _style_data(ws, total_row, len(headers), devise)
    _auto_width(ws, len(headers))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Cotisations_Imprevus.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_assistances_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Assistances"
    devise = get_devise(request)

    headers = ['#', 'Membre', 'Carte ID', 'Montant', 'Motif', 'Date assistance', 'Créé par', 'Date création']
    ws.append(headers)
    _style_header(ws, len(headers))

    assistances = qs_for_tenant(request, CasAssistance.objects.all()).select_related('membre', 'cree_par').order_by('-date_assistance')

    total = Decimal('0')
    for i, a in enumerate(assistances, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=a.membre.get_full_name())
        ws.cell(row=row, column=3, value=a.membre.carte_identite)
        _format_money(ws, row, 4, a.montant)
        ws.cell(row=row, column=5, value=a.motif)
        ws.cell(row=row, column=6, value=a.date_assistance.strftime('%d/%m/%Y') if a.date_assistance else '')
        ws.cell(row=row, column=7, value=a.cree_par.get_full_name() if a.cree_par else '')
        ws.cell(row=row, column=8, value=a.date_creation.strftime('%d/%m/%Y %H:%M') if a.date_creation else '')
        total += a.montant

    total_row = len(assistances) + 2
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    _format_money(ws, total_row, 4, total)
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    _style_data(ws, total_row, len(headers), devise)
    _auto_width(ws, len(headers))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Assistances.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_interets_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Intérêts"
    devise = get_devise(request)

    headers = ['#', 'Membre', 'Crédit #', 'Mois', 'Année', 'Intérêt mensuel', 'Intérêt cumulé', 'Date calcul']
    ws.append(headers)
    _style_header(ws, len(headers))

    interets = qs_for_tenant(request, Interet.objects.all()).select_related('credit__membre')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')
    if date_debut:
        interets = interets.filter(date_calcul__gte=date_debut)
    if date_fin:
        interets = interets.filter(date_calcul__lte=date_fin)
    
    interets = interets.order_by('-annee', '-mois')
    mois_noms = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']

    total_mensuel = Decimal('0')
    total_cumule = Decimal('0')
    for i, it in enumerate(interets, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=it.credit.membre.get_full_name())
        ws.cell(row=row, column=3, value=it.credit.id_credit)
        ws.cell(row=row, column=4, value=mois_noms[it.mois])
        ws.cell(row=row, column=5, value=it.annee)
        _format_money(ws, row, 6, it.interet_mensuel)
        _format_money(ws, row, 7, it.interet_cumule)
        ws.cell(row=row, column=8, value=it.date_calcul.strftime('%d/%m/%Y') if it.date_calcul else '')
        total_mensuel += it.interet_mensuel
        total_cumule += it.interet_cumule

    total_row = len(interets) + 2
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    _format_money(ws, total_row, 6, total_mensuel)
    _format_money(ws, total_row, 7, total_cumule)
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    _style_data(ws, total_row, len(headers), devise)
    _auto_width(ws, len(headers))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Interets.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_remboursements_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Remboursements"
    devise = get_devise(request)

    headers = ['#', 'Membre', 'Crédit #', 'Montant', 'Reste', 'Mode', 'Date']
    ws.append(headers)
    _style_header(ws, len(headers))

    remboursements = qs_for_tenant(request, Remboursement.objects.all()).select_related('credit__membre').order_by('-date_remboursement')

    total_principal = Decimal('0')
    for i, r in enumerate(remboursements, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=r.credit.membre.get_full_name())
        ws.cell(row=row, column=3, value=r.credit.id_credit)
        _format_money(ws, row, 4, r.montant_rembourse)
        _format_money(ws, row, 5, r.credit.get_reste_a_rembourser())
        ws.cell(row=row, column=6, value=r.get_mode_paiement_display())
        ws.cell(row=row, column=7, value=r.date_remboursement.strftime('%d/%m/%Y') if r.date_remboursement else '')
        total_principal += r.montant_rembourse

    total_row = len(remboursements) + 2
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    _format_money(ws, total_row, 4, total_principal)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Remboursements.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_rapport_mensuel_excel(request, mois, annee):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Rapport {mois}_{annee}"
    devise = get_devise(request)

    tenant = get_tenant(request)
    association = Association.get_association(tenant)
    nom_assoc = association.nom if association else 'Association'

    mois_noms = ['', 'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin',
                 'Juillet', 'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre']
    mois_nom = mois_noms[mois]

    base_cotisations = qs_for_tenant(request, Cotisation.objects.all())
    base_remboursements = qs_for_tenant(request, Remboursement.objects.all())
    base_credits = qs_for_tenant(request, Credit.objects.all())
    base_interets = qs_for_tenant(request, Interet.objects.all())
    cotisations_mois = base_cotisations.filter(mois=mois, annee=annee).select_related('membre')
    remboursements_mois = base_remboursements.filter(
        date_remboursement__month=mois, date_remboursement__year=annee
    ).select_related('credit__membre')
    interets_mois = base_interets.filter(
        date_remboursement__month=mois, date_remboursement__year=annee
    ).select_related('credit__membre')
    credits_mois = base_credits.filter(
        statut='APPROUVE', date_accord__month=mois, date_accord__year=annee
    ).select_related('membre')

    membres_actifs = qs_for_tenant(request, Membre.objects.filter(est_actif=True))

    headers = ['Membre', f'Cotisation ({devise})', f'Remboursement ({devise})',
               f'Intérêt ({devise})', f'Crédit ({devise})', f'Montant généré ({devise})']
    ws.append(headers)
    _style_header(ws, len(headers))

    cotisations_dict = {}
    for c in cotisations_mois:
        key = c.membre_id
        cotisations_dict[key] = cotisations_dict.get(key, Decimal('0')) + c.montant

    remb_dict = {}
    for r in remboursements_mois:
        key = r.credit.membre_id
        remb_dict[key] = remb_dict.get(key, Decimal('0')) + r.montant_rembourse

    interets_dict = {}
    for it in interets_mois:
        key = it.credit.membre_id
        interets_dict[key] = interets_dict.get(key, Decimal('0')) + it.interet_rembourse

    credits_dict = {}
    for cr in credits_mois:
        key = cr.membre_id
        credits_dict[key] = credits_dict.get(key, Decimal('0')) + cr.montant_accorde

    all_member_ids = set()
    all_member_ids.update(cotisations_dict.keys())
    all_member_ids.update(remb_dict.keys())
    all_member_ids.update(interets_dict.keys())
    all_member_ids.update(credits_dict.keys())
    for m in membres_actifs:
        all_member_ids.add(m.pk)

    membres_map = {m.pk: m for m in membres_actifs}
    for c in cotisations_mois:
        if c.membre_id not in membres_map and c.membre:
            membres_map[c.membre_id] = c.membre
    for r in remboursements_mois:
        if r.credit.membre_id not in membres_map:
            membres_map[r.credit.membre_id] = r.credit.membre
    for cr in credits_mois:
        if cr.membre_id not in membres_map:
            membres_map[cr.membre_id] = cr.membre

    row = 2
    total_cot = Decimal('0')
    total_remb = Decimal('0')
    total_int = Decimal('0')
    total_cred = Decimal('0')
    total_gen = Decimal('0')

    for mid in sorted(all_member_ids):
        membre = membres_map.get(mid)
        if not membre:
            continue
        cot_val = cotisations_dict.get(mid, Decimal('0'))
        remb_val = remb_dict.get(mid, Decimal('0'))
        int_val = interets_dict.get(mid, Decimal('0'))
        cred_val = credits_dict.get(mid, Decimal('0'))
        gen_val = cot_val + remb_val + int_val - cred_val

        ws.cell(row=row, column=1, value=membre.get_full_name())
        _format_money(ws, row, 2, cot_val)
        _format_money(ws, row, 3, remb_val)
        _format_money(ws, row, 4, int_val)
        _format_money(ws, row, 5, cred_val)

        formula = f"=ROUND(B{row}+C{row}+D{row}-E{row},3)"
        ws.cell(row=row, column=6).value = formula
        ws.cell(row=row, column=6).number_format = '#,##0.000'

        total_cot += cot_val
        total_remb += remb_val
        total_int += int_val
        total_cred += cred_val
        total_gen += gen_val
        row += 1

    total_row = row
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    _format_money(ws, total_row, 2, total_cot)
    _format_money(ws, total_row, 3, total_remb)
    _format_money(ws, total_row, 4, total_int)
    _format_money(ws, total_row, 5, total_cred)
    ws.cell(row=total_row, column=6, value=f"=ROUND(B{total_row}+C{total_row}+D{total_row}-E{total_row},3)")
    ws.cell(row=total_row, column=6).number_format = '#,##0.000'
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    _style_data(ws, total_row, len(headers), devise)
    _auto_width(ws, len(headers))

    ws.cell(row=1, column=1).value = f"Rapport Mensuel - {nom_assoc} - {mois_nom} {annee}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal='center')

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Rapport_Mensuel_{mois_nom}_{annee}.xlsx"'
    wb.save(response)
    return response


@login_required
@user_passes_test(is_admin_or_gestionnaire)
def export_penalites_excel(request):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pénalités"
    devise = get_devise(request)

    from .models import PenaliteCotisation
    headers = ['#', 'Membre', 'Période', 'Cotisation', 'Taux', 'Pénalité mensuelle', 'Pénalité cumulée', 'Statut', 'Date paiement']
    ws.append(headers)
    _style_header(ws, len(headers))

    penalites = qs_for_tenant(request, PenaliteCotisation.objects.filter(exclu=False)).filter(membre__est_actif=True).select_related('membre')
    date_debut = request.GET.get('date_debut')
    date_fin = request.GET.get('date_fin')
    if date_debut:
        # Comme pénalité peut ne pas avoir de date_paiement, on filtre sur date_paiement si payée ou date_creation/autre, 
        # mais vu qu'il n'y a peut être pas date_creation on va filtrer sur date_paiement pour l'instant.
        penalites = penalites.filter(date_paiement__gte=date_debut)
    if date_fin:
        penalites = penalites.filter(date_paiement__lte=date_fin)
    
    penalites = penalites.order_by('-annee', '-mois')

    total_mensuel = Decimal('0')
    total_cumule = Decimal('0')
    for i, p in enumerate(penalites, 1):
        row = i + 1
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=p.membre.get_full_name())
        periode = f"{p.get_mois_nom()} {p.annee}" if p.mois else "-"
        ws.cell(row=row, column=3, value=periode)
        _format_money(ws, row, 4, p.montant_cotisation)
        ws.cell(row=row, column=5, value=f"{p.taux_penalite}%")
        _format_money(ws, row, 6, p.penalite_mensuelle)
        _format_money(ws, row, 7, p.penalite_cumulee)
        ws.cell(row=row, column=8, value='Payée' if p.est_payee else 'En attente')
        ws.cell(row=row, column=9, value=p.date_paiement.strftime('%d/%m/%Y') if p.date_paiement else '')
        
        total_mensuel += p.penalite_mensuelle
        total_cumule += p.penalite_cumulee

    total_row = len(penalites) + 2
    ws.cell(row=total_row, column=1, value='TOTAL')
    ws.cell(row=total_row, column=1).font = Font(bold=True)
    _format_money(ws, total_row, 6, total_mensuel)
    _format_money(ws, total_row, 7, total_cumule)
    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).font = Font(bold=True)

    _style_data(ws, total_row, len(headers), devise)
    _auto_width(ws, len(headers))

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Penalites.xlsx"'
    wb.save(response)
    return response


def generer_backup_tenant(tenant, request=None):
    from datetime import date as date_func
    wb = Workbook()
    devise = 'FBU'
    if tenant:
        assoc = Association.get_association(tenant)
        devise = assoc.devise if assoc and assoc.devise else 'FBU'
    now = datetime.now().strftime('%Y%m%d_%H%M%S')

    def qs(qs_model):
        qs = qs_model
        if tenant:
            return qs.filter(tenant=tenant)
        return qs

    def write_sheet(ws, title, headers, rows_data):
        ws.title = title
        ws.append(headers)
        hdr_font = Font(bold=True, color='FFFFFF', size=11)
        hdr_fill = PatternFill(start_color='4a90e2', end_color='4a90e2', fill_type='solid')
        hdr_align = Alignment(horizontal='center', vertical='center')
        thin = Border(left=Side(style='thin'), right=Side(style='thin'),
                      top=Side(style='thin'), bottom=Side(style='thin'))
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.font = hdr_font; c.fill = hdr_fill; c.alignment = hdr_align; c.border = thin
        for r_idx, row_vals in enumerate(rows_data, 2):
            for c_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.font = Font(size=10); cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = thin
        for col in range(1, len(headers) + 1):
            max_len = max((len(str(ws.cell(row=r, column=col).value or '')) for r in range(1, len(rows_data) + 2)), default=0)
            ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 40)

    # 1. Membres
    ws = wb.active
    rows = []
    membres = Membre.objects.filter(is_superuser=False)
    if tenant: membres = membres.filter(tenant=tenant)
    for m in membres.order_by('nom', 'prenom'):
        rows.append([m.get_full_name(), m.carte_identite or '', m.email or '',
                     m.telephone or '', m.adresse or '', 'Actif' if m.est_actif else 'Inactif',
                     m.date_inscription.strftime('%d/%m/%Y') if m.date_inscription else ''])
    write_sheet(ws, 'Membres', ['Nom','Carte ID','Email','Téléphone','Adresse','Statut','Date inscription'], rows)

    # 2. Cotisations
    ws = wb.create_sheet()
    rows = []
    for c in qs(Cotisation.objects.filter(statut='PAYE')).select_related('membre').order_by('-annee','-mois'):
        rows.append([c.membre.get_full_name(), c.get_mois_nom(), c.annee, float(c.montant),
                     c.date_paiement.strftime('%d/%m/%Y') if c.date_paiement else ''])
    write_sheet(ws, 'Cotisations', ['Membre','Mois','Année','Montant','Date paiement'], rows)

    # 3. Imprévus
    ws = wb.create_sheet()
    rows = []
    for c in qs(CotisationImprevu.objects.filter(statut='PAYE')).select_related('membre').order_by('-annee','-mois'):
        rows.append([c.membre.get_full_name(), c.get_mois_nom(), c.annee, float(c.montant),
                     c.date_paiement.strftime('%d/%m/%Y') if c.date_paiement else ''])
    write_sheet(ws, 'Imprévus', ['Membre','Mois','Année','Montant','Date paiement'], rows)

    # 4. Crédits
    ws = wb.create_sheet()
    rows = []
    for c in qs(Credit.objects.all()).select_related('membre').order_by('-date_demande'):
        rows.append([c.membre.get_full_name(), float(c.montant_demande), float(c.montant_accorde),
                     c.get_statut_display(), c.date_demande.strftime('%d/%m/%Y'),
                     c.date_accord.strftime('%d/%m/%Y') if c.date_accord else ''])
    write_sheet(ws, 'Crédits', ['Membre','Montant demandé','Montant accordé','Statut','Date demande','Date accord'], rows)

    # 5. Remboursements
    ws = wb.create_sheet()
    rows = []
    for r in qs(Remboursement.objects.all()).select_related('credit__membre').order_by('-date_remboursement'):
        rows.append([r.credit.membre.get_full_name(), r.credit.id_credit, float(r.montant_rembourse),
                     r.get_mode_paiement_display() if hasattr(r,'get_mode_paiement_display') else '',
                     r.date_remboursement.strftime('%d/%m/%Y') if r.date_remboursement else ''])
    write_sheet(ws, 'Remboursements', ['Membre','Crédit #','Montant','Mode','Date'], rows)

    # 6. Intérêts
    ws = wb.create_sheet()
    rows = []
    for it in qs(Interet.objects.all()).select_related('credit__membre').order_by('-annee','-mois'):
        rows.append([it.credit.membre.get_full_name(), it.credit.id_credit, it.get_mois_nom(), it.annee,
                     float(it.interet_mensuel), float(it.interet_cumule),
                     it.date_calcul.strftime('%d/%m/%Y') if it.date_calcul else ''])
    write_sheet(ws, 'Intérêts', ['Membre','Crédit #','Mois','Année','Intérêt mensuel','Intérêt cumulé','Date calcul'], rows)

    # 7. Pénalités
    ws = wb.create_sheet()
    rows = []
    for p in qs(PenaliteCotisation.objects.filter(exclu=False)).select_related('membre').order_by('-annee','-mois'):
        rows.append([p.membre.get_full_name(), f"{p.get_mois_nom()} {p.annee}" if p.mois else '',
                     float(p.montant_cotisation), f"{p.taux_penalite}%",
                     float(p.penalite_mensuelle), float(p.penalite_cumulee),
                     'Payée' if p.est_payee else 'En attente',
                     p.date_paiement.strftime('%d/%m/%Y') if p.date_paiement else ''])
    write_sheet(ws, 'Pénalités', ['Membre','Période','Cotisation','Taux','Pénalité mensuelle','Pénalité cumulée','Statut','Date paiement'], rows)

    # 8. Assistances
    ws = wb.create_sheet()
    rows = []
    for a in qs(CasAssistance.objects.all()).select_related('membre').order_by('-date_assistance'):
        rows.append([a.membre.get_full_name(), float(a.montant), a.motif,
                     a.date_assistance.strftime('%d/%m/%Y') if a.date_assistance else ''])
    write_sheet(ws, 'Assistances', ['Membre','Montant','Motif','Date'], rows)

    # 9. Suggestions
    ws = wb.create_sheet()
    rows = []
    for s in qs(Suggestion.objects.all()).select_related('membre').order_by('-date_soumission'):
        rows.append([s.membre.get_full_name(), s.titre, s.description or '',
                     s.date_soumission.strftime('%d/%m/%Y') if s.date_soumission else '',
                     s.get_statut_display()])
    write_sheet(ws, 'Suggestions', ['Membre','Titre','Description','Date','Statut'], rows)

    nom_tenant = tenant.nom.replace(' ', '_') if tenant else 'TOUTES'
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer, nom_tenant, now


def export_backup_complet(request):
    tenant = get_tenant(request)
    buffer, nom_tenant, now = generer_backup_tenant(tenant, request)
    response = HttpResponse(buffer.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Backup_{nom_tenant}_{now}.xlsx"'
    return response
